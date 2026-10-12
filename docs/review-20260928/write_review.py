from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

out=Path('docs/review-20260928/De-xuat-nghiep-vu-can-duyet.docx')
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.8)
normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.line_spacing=1.1
for name,size in [('Title',19),('Heading 1',13)]:
 s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(5)
doc.add_paragraph('Các điểm nghiệp vụ MotoRent cần duyệt', 'Title')
doc.add_paragraph('Gửi chủ dự án • Rà soát ngày 28/09/2026')
doc.add_paragraph('Tôi ghi lại các điểm dưới đây từ mã nguồn hiện tại để bạn quyết định trước khi thay đổi chức năng đã có. Đây là đề xuất chưa áp dụng, tách khỏi các sửa lỗi ngày thuê, bản đồ, cọc và dữ liệu địa điểm bạn đã yêu cầu.')
doc.add_paragraph('Hai CCCD cần thuộc hai người khác nhau', 'Heading 1')
doc.add_paragraph('Hiện tại, tạo đơn kiểm tra số lượng người và định dạng từng CCCD nhưng chưa đối chiếu trùng số giữa hai người. Vì phương án hai CCCD được miễn cọc tài sản, khách có thể nhập lặp cùng giấy tờ để nhận mức cọc 0.')
doc.add_paragraph('Đề xuất: từ chối hai số giấy tờ trùng nhau sau khi chuẩn hóa khoảng trắng và chữ hoa; yêu cầu nhân viên đối chiếu giấy tờ khi giao xe. Bạn xác nhận có áp dụng quy tắc này không.')
doc.add_paragraph('Điều kiện bàn giao và thu cọc tài sản', 'Heading 1')
doc.add_paragraph('Hàm bàn giao hiện cho chuyển đơn đã xác nhận sang đang thuê mà chưa bắt buộc có biên bản tình trạng xe hoặc đã ghi nhận đủ cọc tài sản. Hàm cũng chưa kiểm tra mốc nhận xe, nên có thể bàn giao sớm dù lịch đặt trên hệ thống chưa bắt đầu.')
doc.add_paragraph('Đề xuất: trước bàn giao, kiểm tra giấy tờ, biên bản, khoản cọc theo đơn và lịch trống thực tế. Trường hợp cho nhận sớm hoặc miễn cọc cần quyền quản trị và lý do lưu lại. Bạn chọn có bắt buộc thu đủ cọc và cho phép nhận sớm trong trường hợp nào.')
doc.add_paragraph('Chỉnh sửa thông tin xe điện', 'Heading 1')
doc.add_paragraph('Danh mục có xe điện dùng dung tích 0 cc, nhưng API lưu xe của admin chỉ nhận từ 40 đến 2000 cc. Vì vậy, chỉnh sửa một xe điện hiện có sẽ bị từ chối dù không thay dung tích.')
doc.add_paragraph('Đề xuất: cho phép 0 cc riêng đối với loại xe điện trong database; giữ kiểm tra dung tích dương cho xe động cơ đốt trong. Chưa đổi quy tắc này để chờ bạn duyệt.')
doc.add_paragraph('Căn cứ kiểm tra', 'Heading 1')
doc.add_paragraph('RentalOrderService.createOrder và handOverOrder; OperationsService.save mục motorbikes. Các nhận xét là kết quả đọc mã; chưa sửa các nhánh nghiệp vụ nêu trên. Bộ kiểm thử 29 bài của đợt sửa trước đạt, nhưng không chứng minh các quy tắc đề xuất đã được triển khai.')
doc.core_properties.title='Các điểm nghiệp vụ MotoRent cần duyệt'
doc.core_properties.author='Codex'
doc.save(out)
print(out.resolve())

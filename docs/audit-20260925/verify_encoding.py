from pathlib import Path
from html.parser import HTMLParser
from zipfile import ZipFile
from docx import Document
import json,re
root=Path(__file__).resolve().parents[2]; out=Path(__file__).parent
class Tags(HTMLParser):
 def __init__(self): super().__init__(convert_charrefs=False); self.tokens=[]
 def handle_starttag(self,tag,attrs): self.tokens.append(('start',tag,[k for k,v in attrs]))
 def handle_endtag(self,tag): self.tokens.append(('end',tag))
paths=json.loads((out/'encoding-files.json').read_text('utf-8'))
for name in paths:
 old=(out/'originals'/name).read_bytes().decode('utf-8')
 new=(root/name).read_bytes().decode('utf-8')
 a=Tags(); b=Tags(); a.feed(old); b.feed(new)
 assert a.tokens==b.tokens,name+' structure changed'
 assert re.sub(r'[^\x00-\x7f]','',old)==re.sub(r'[^\x00-\x7f]','',new),name+' ASCII logic changed'
 assert '\ufffd' not in new
 assert not any(128<=ord(c)<=159 for c in new)
 assert (root/'target/classes/templates/customer'/Path(name).name).read_bytes()==(root/name).read_bytes()
 print('PASS UTF-8, HTML structure, ASCII code unchanged, compiled resource matches: '+name)
report=out/'Bao-cao-ra-soat-MotoRent.docx'
with ZipFile(report) as z: assert z.testzip() is None
doc=Document(report)
assert len(doc.paragraphs)>50
assert all('\ufffd' not in p.text for p in doc.paragraphs)
print('PASS DOCX archive and Vietnamese text integrity (visual rendering unavailable)')

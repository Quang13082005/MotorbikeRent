from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT=Path(__file__).parent
d=Document(); s=d.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(1.8); s.left_margin=s.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(10.5 if name=='Normal' else 22 if name=='Title' else 16 if name=='Heading 1' else 12)
 st.paragraph_format.space_after=Pt(7)
 st.paragraph_format.line_spacing=1.12
d.core_properties.title='Rà soát và đề xuất hoàn thiện MotoRent'
d.core_properties.author='Codex'
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def finding(title,body,source):
 h(title); p(body)
 q=d.add_paragraph('Vị trí: '+source); q.paragraph_format.space_after=Pt(9)
 for r in q.runs: r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string('555555')
d.add_heading('Rà soát và đề xuất hoàn thiện MotoRent',0)
p('Gửi chủ dự án cho thuê xe máy tại Đà Nẵng • Ngày 25 tháng 9 năm 2026')
p('Tôi đề nghị ưu tiên sửa việc xác định người đăng nhập, phân quyền đơn thuê và xác nhận thanh toán trước khi nối các màn hình còn dùng dữ liệu mẫu vào MySQL. Các thay đổi nghiệp vụ dưới đây cần bạn xem và duyệt theo yêu cầu ban đầu. Chưa có bảng hay dữ liệu MySQL nào bị sửa hoặc xóa trong đợt rà soát này.')
h('Kết nối và phần đã sửa')
p('Kết nối trực tiếp bằng JDBC thành công tới motorent_db tại localhost:3306, dùng cấu hình sẵn của dự án. Không cần MCP để thực hiện công việc này. Database có 21 bảng, dùng utf8mb4 và collation utf8mb4_0900_ai_ci. Chỉ truy vấn số lượng bản ghi và quan hệ khóa ngoại, không đưa mật khẩu hay token vào báo cáo.')
p('Đã sửa các chuỗi tiếng Việt bị giải mã sai trong profile.html và rental-history.html, lưu đúng UTF-8. Đã giữ bản gốc trong docs/audit-20260925/originals. Bản sửa giới hạn ở chữ hiển thị và chuỗi văn bản; chưa đổi controller, API hay quy tắc nghiệp vụ.')
finding('Ưu tiên 1 Xác định đúng khách hàng',
'Trang hồ sơ, đơn đang thuê, lịch sử, API tạo đơn và API đánh giá lấy CUSTOMER đầu tiên bằng findAll().stream().findFirst(). Vì vậy thông tin và đơn có thể thuộc người khác dù khách đã đăng nhập. Đề xuất lấy userId từ principal đã xác thực, truy vấn theo chủ đơn; tài khoản thiếu dữ liệu phải hiển thị trạng thái trống, không dùng khách mẫu thay thế.',
'CustomerPageController.java:97,124,148; RentalOrderApiController.java:29,40; ReviewApiController.java:25.')
finding('Ưu tiên 1 Phân quyền thao tác đơn',
'API xem chi tiết và hủy đơn chưa kiểm tra chủ sở hữu. Các API xác nhận, từ chối, bàn giao, hoàn tất chưa chặn theo vai trò và lấy STAFF đầu tiên làm người xử lý. SecurityConfig chỉ giới hạn /staff/** và /admin/**; /api/rental-orders/** rơi vào yêu cầu đăng nhập chung. Đề xuất CUSTOMER chỉ thao tác đơn của mình; STAFF và ADMIN xử lý theo quyền, ghi nhận đúng người thao tác. Kiểm thử phải có hai khách hàng khác nhau và truy cập API trực tiếp.',
'RentalOrderApiController.java:50–119; SecurityConfig.java; RentalOrderService.java.')
page('Thanh toán và vòng đời thuê xe')
finding('Ưu tiên 1 Không xác nhận thanh toán bằng URL',
'paymentResult() đổi đơn từ PENDING_PAYMENT sang PENDING chỉ dựa vào code=00, cancel=false và orderId trên URL công khai. Đây là đầu vào người dùng có thể thay đổi, chưa phải bằng chứng đã nhận tiền. Đề xuất trang kết quả chỉ đọc trạng thái trong database; xác nhận thanh toán từ webhook đã xác minh hoặc truy vấn PayOS ở máy chủ, đối chiếu số tiền và mã giao dịch đã lưu.',
'CustomerPageController.java:159–179; SecurityConfig.java; PaymentController.java:35.')
finding('Ưu tiên 1 Lưu giao dịch và tính tiền tại máy chủ',
'API tạo link nhận amount từ người gọi; PaymentService chưa lưu bản ghi payments. Webhook có xác minh chữ ký nhưng tìm đơn bằng nội dung mô tả, chưa đối chiếu amount với khoản phải thu đã lưu. Đề xuất lưu mã PayOS duy nhất cho mỗi lần thanh toán, số tiền kỳ vọng, loại khoản thu và trạng thái. Dùng giao dịch database cùng cơ chế xử lý lặp an toàn để webhook gửi lại không ghi nhận tiền hai lần. Không nhận số tiền quyết định từ trình duyệt.',
'PaymentService.java:createPaymentLink; PaymentController.java:createPaymentLink,handlePayOSWebhook; entity/Payment.java.')
finding('Ưu tiên 2 Giá thuê và giữ chỗ',
'createOrder() suy ra phí giao xe từ depositAmount khách gửi lên; ChronoUnit.HOURS bỏ phần phút trước khi làm tròn ngày. Ví dụ 24 giờ 30 phút có thể bị tính một ngày dù chính sách là làm tròn mỗi 24 giờ. Kiểm tra trùng lịch rồi lưu chưa khóa cạnh tranh; chưa chặn xe bảo dưỡng bằng trạng thái xe. Đề xuất chốt quy tắc làm tròn, biểu phí giao nhận và kiểm tra khả dụng ở máy chủ, khóa phù hợp khi đặt cùng xe.',
'RentalOrderService.java:createOrder; RentalOrderRepository.java:existsOverlappingOrder.')
finding('Ưu tiên 2 Hết hạn cọc và hoàn tất chuyến đi',
'Đơn hết hạn sau ba phút tính từ createdAt, còn link PayOS hết hạn ba phút tính từ lần tạo link; lần thanh toán lại có thể lệch hạn giữ chỗ. Hàm rejectOrder(orderId,staffId,reason) đang được API gọi không kiểm tra trạng thái như overload còn lại. completeOrder() chuyển hoàn tất mà chưa ghi nhận đối soát tiền còn thiếu hoặc tình trạng xe. Đề xuất thống nhất hạn thanh toán, xử lý webhook đến muộn và chốt điều kiện từ chối, trả xe, hoàn cọc trước khi sửa.',
'RentalOrderService.java:cleanupExpiredOrders,rejectOrder,completeOrder; PaymentService.java.')
page('Các màn hình cần thay dữ liệu mẫu')
h('Customer cần hoàn thiện trước')
p('Hồ sơ: tên, email, ngày sinh, giới tính, địa chỉ và ảnh có giá trị mẫu. saveProfile() chỉ hiện thông báo; tải ảnh chỉ alert; form đổi mật khẩu cũng chỉ alert rồi reset. Cần API đọc và lưu hồ sơ theo người đăng nhập, kiểm tra mật khẩu hiện tại, xác nhận mật khẩu mới và upload có lưu trữ thật. Các nhãn đã xác minh chỉ hiện khi có trạng thái xác minh tương ứng trong database.')
p('Lịch sử: số lượng ở đầu trang đã lấy từ đơn thuê nhưng chân trang cố định “1 - 4” trên tổng “5”; nút trang 2 và nút đánh giá chưa nối chức năng. Bộ lọc thời gian chưa áp dụng dữ liệu. Ảnh ghép thêm dấu / trước imagePrimary có thể làm hỏng URL tuyệt đối. Cần phân trang và lọc thực, đường dẫn ảnh hợp lệ, trạng thái trống, liên kết đánh giá đúng đơn và thuê lại đúng xe.')
p('Trang chủ: form tìm theo ngày đang alert rồi cuộn tới bảng giá; cần truyền ngày nhận/trả tới truy vấn xe còn trống. Trang chi tiết xe và danh sách đã dùng service/database, nên giữ luồng này và bổ sung dữ liệu còn thiếu thay vì viết lại toàn bộ.')
p('Gia hạn, đánh giá, thanh toán và sửa hồ sơ: nhiều route chỉ trả template mà không cấp model cần thiết. review.html và rental-extension.html tham chiếu rental nhưng controller chưa nạp; form đang gửi GET tới chi tiết đơn, chưa lưu yêu cầu. Cần nối API/database, kiểm tra quyền và trạng thái, lưu yêu cầu gia hạn, kiểm tra trùng lịch trước khi chấp nhận. ReviewService đã có kiểm tra chủ đơn, đơn hoàn tất và đánh giá trùng; giữ các kiểm tra này khi sửa nguồn customerId.')
h('Admin và Staff')
p('Admin: dashboard, báo cáo, xe, khách, nhân viên, đơn thuê, khuyến mãi, cài đặt và bảo dưỡng còn dữ liệu mẫu. operations.js chứa mảng dữ liệu và thao tác trên bộ nhớ trình duyệt; tải lại sẽ mất thay đổi. Cần API theo quyền ADMIN, truy vấn database, kiểm tra dữ liệu nhập, lưu thật và xuất báo cáo từ giao dịch đã đối soát. Các nhãn, danh sách trạng thái và cấu trúc giao diện có thể là hằng số; dữ liệu khách, xe, đơn và doanh thu thì không.')
p('Staff: dashboard và danh sách đơn đã đọc RentalOrderService; hồ sơ gọi /api/auth/me. Giữ các kết nối này. Hoàn thiện phân quyền API, người thực hiện và kiểm tra trạng thái; bảo dưỡng cần lưu maintenance_records và cập nhật khả dụng xe trong cùng giao dịch.')
h('Đăng nhập và dữ liệu khởi tạo')
p('DatabaseSeeder tạo tài khoản demo cả khi không bật profile preview; nên giới hạn seed cho môi trường thử nghiệm. Logout đổi isUsed của token trong database nhưng JwtAuthenticationFilter chỉ kiểm tra chữ ký và hạn JWT; cần áp dụng trạng thái thu hồi và trạng thái tài khoản khi xác thực. SMTP hiện là cấu hình mẫu, nên luồng quên mật khẩu cần cấu hình gửi mail và kiểm thử riêng trước nghiệm thu.')
page('Đề xuất giữ và tinh gọn database')
p('Bảng rỗng không đồng nghĩa với bảng thừa. Số lượng dưới đây là ảnh chụp tại thời điểm truy vấn và có thể thay đổi. Chưa đề nghị DROP TABLE ngay cho bất kỳ bảng nào.')
h('Giữ cho nghiệp vụ đang dùng')
p('users 7; token_authentication 43; otp_tokens 0; motorbikes 11; motorbike_brands 4; motorbike_categories 4; motorbike_images 5; rental_orders 2; reviews 0. Đây là dữ liệu nền của tài khoản, xác thực, danh mục xe, đặt thuê và đánh giá.')
h('Giữ và nối vào chức năng thực tế')
p('payments 0: nên giữ dù dùng PayOS. PayOS xử lý tiền, còn bảng nội bộ liên kết giao dịch với đơn, theo dõi cọc và phần còn lại, đối soát, xử lý webhook lặp và hoàn tiền. Cần bổ sung ràng buộc mã giao dịch duy nhất và thông tin khoản thu phù hợp.')
p('maintenance_records 0; rental_contracts 0; rental_extensions 0; user_documents 0; notifications 0; promotions 0. Tương ứng bảo dưỡng, hợp đồng, gia hạn, xác minh giấy tờ, thông báo và khuyến mãi. Cần hoàn thiện luồng sử dụng; chỉ cân nhắc bỏ khi bạn quyết định bỏ luôn chức năng tương ứng.')
h('Một cửa hàng và bảng stations')
p('stations có 2 bản ghi; motorbikes.station_id đang có khóa ngoại tới bảng này và MotorbikeMapper đọc station.name. Xóa ngay sẽ ảnh hưởng dữ liệu xe hoặc lỗi ánh xạ. Đề xuất giai đoạn đầu giữ bảng và chuẩn hóa về cửa hàng thật sau khi xác nhận địa chỉ; nếu muốn bỏ mô hình trạm, chuyển thông tin cửa hàng sang shop_settings, chuyển các tham chiếu và kiểm thử rồi mới loại bỏ khóa ngoại, cột và bảng theo migration có bản sao lưu.')
h('Các bảng cần xác minh nguồn gốc')
p('shop_settings 0: phù hợp địa chỉ, giờ mở cửa và cấu hình một cửa hàng. deposit_attempts 0: chưa thấy mã Java dùng; có thể trùng vai trò lưu lần thanh toán với payments. guest_bookings 2 và travel_places 28: chưa thấy mã Java dùng nhưng đã có dữ liệu; cần xác minh có thuộc phiên bản khác hay công cụ ngoài dự án. Chưa đủ bằng chứng để xóa ba bảng này hoặc kết luận dữ liệu là không cần thiết.')
h('Điều kiện trước khi xóa bảng')
p('Chốt phạm vi chức năng; kiểm tra tham chiếu ở mã nguồn, khóa ngoại, view, trigger, event và ứng dụng khác; xuất bản sao lưu và thử khôi phục; chuyển dữ liệu cần giữ; chạy migration trên database thử nghiệm rồi kiểm thử tất cả vai trò. JPA đang đặt ddl-auto=update, nên cần quản lý migration để tránh bảng hoặc cột đã bỏ được tạo lại.')
page('Phạm vi duyệt và tiêu chí nghiệm thu')
h('Các quyết định cần bạn xem')
p('1. Duyệt sửa xác thực, quyền sở hữu đơn và phân quyền API trước khi nối giao diện. Chỉ đọc mã và truy vấn database trong lần rà soát này, chưa thử khai thác hoặc thay đổi trạng thái đơn thật.')
p('2. Duyệt mô hình lưu giao dịch PayOS trong payments, tiền cọc và hạn giữ chỗ thống nhất. Cần chốt quy tắc hoàn cọc khi hủy, từ chối, thanh toán đến muộn và khoản còn thu khi trả xe.')
p('3. Xác nhận cách tính ngày thuê lẻ, phí giao nhận, xử lý trả muộn và gia hạn. Đề xuất mỗi 24 giờ làm tròn lên nếu phù hợp cách kinh doanh thực tế; chưa áp dụng chính sách này.')
p('4. Với một cửa hàng, đề xuất giữ stations trong giai đoạn chuyển đổi rồi cân nhắc chuyển sang shop_settings. Cần xác nhận cửa hàng thực tế và nguồn gốc guest_bookings, travel_places trước khi đề nghị xóa dữ liệu.')
h('Thứ tự triển khai sau khi duyệt')
p('Customer: danh tính và quyền truy cập → thanh toán → hồ sơ và đổi mật khẩu → tìm xe theo lịch → đơn thuê và lịch sử → đánh giá và gia hạn. Nghiệm thu customer trước khi tiếp tục admin rồi staff theo yêu cầu của bạn.')
p('Admin: danh mục xe, khách và nhân viên → đơn thuê → bảo dưỡng → giao dịch và báo cáo → cài đặt và khuyến mãi. Staff: giữ các màn hình đã đọc database, hoàn thiện xử lý đơn và bảo dưỡng dựa trên backend đã thống nhất.')
h('Kiểm thử cần có trước nghiệm thu')
p('Hai khách không xem hoặc sửa đơn của nhau; CUSTOMER không gọi được thao tác nhân viên. Hồ sơ lưu lại sau tải trang; mật khẩu cũ sai bị từ chối. Hai yêu cầu đồng thời không đặt trùng xe; xe bảo dưỡng không nhận đơn. Bộ lọc và phân trang hoạt động với 0, 1 và nhiều đơn.')
p('URL kết quả không thay đổi thanh toán; amount từ trình duyệt không quyết định khoản thu. Webhook sai chữ ký, sai tiền, gửi lặp và đến muộn được xử lý đúng. Báo cáo không tính cọc giữ hộ thành doanh thu tùy chính sách đã duyệt; đối chiếu tổng với giao dịch database. Kiểm thử dữ liệu riêng, không cần chuyển tiền thật để kiểm tra nghiệp vụ cốt lõi.')
h('Trạng thái kiểm tra của đợt này')
p('Đã kiểm tra kết nối MySQL và khóa ngoại, rà soát mã các luồng chính và dữ liệu mẫu trong giao diện. Hai file sửa UTF-8 giữ nguyên cấu trúc HTML và mã ASCII; bản trong target/classes khớp nguồn. Maven test chưa chạy được do test cũ không khớp controller: CustomerPaymentResultRouteTests dòng 24 thiếu tham số RentalOrderRepository; OperationsPageControllerTests dòng 39 và 48 gọi handover() và returnMotorbike() không còn tồn tại. Cần cập nhật test theo luồng nghiệp vụ đã duyệt; chưa nghiệm thu toàn hệ thống.')
p('Các đường dẫn Java trong báo cáo thuộc src/main/java/com/motorrent/motorbikerental; giao diện thuộc src/main/resources/templates; operations.js thuộc src/main/resources/static/js. Báo cáo là đề xuất sửa, không phải xác nhận đã hoàn thiện tất cả vai trò.')
out=OUT/'Bao-cao-ra-soat-MotoRent.docx'; d.save(out)
print(out)

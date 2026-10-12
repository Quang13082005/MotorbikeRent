from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path

out=Path('docs/implementation-20260925/Bao-cao-bo-sung-du-lieu-cu.docx')
doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.75)
section.left_margin=section.right_margin=Inches(.85)
for name in ['Normal','Title','Heading 1','Heading 2']:
    style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(21)
doc.styles['Heading 1'].font.size=Pt(14)
doc.add_paragraph('Dữ liệu MotoRent cần xác nhận trước khi sửa',style='Title')
doc.add_paragraph('Báo cáo bổ sung ngày 26 tháng 9 năm 2026')
doc.add_paragraph('Tôi đã đối chiếu trực tiếp database motorent_db bằng truy vấn chỉ đọc. Các điểm dưới đây cần chủ cửa hàng xác nhận để tránh sửa sai lịch sử thuê xe. Mã nguồn mới và kiểm thử hiện chạy trên database bản sao motorent_verify_20260926; chưa sửa hoặc xóa dữ liệu trong database chính.')
sections=[
('Xe chưa có biển số', 'Cả 11 xe trong database chính chưa có biển số. Không thể tự sinh biển số rồi coi đó là xe thực tế. Đề xuất bổ sung biển số và ảnh thực tế trong trang quản trị đội xe. Biển số TEST-CLONE-001 chỉ là dữ liệu kiểm thử ở bản sao, không dùng cho vận hành.'),
('Trạng thái xe không khớp đơn thuê', 'Có 3 xe mang trạng thái RENTED nhưng không có đơn tương ứng ở trạng thái đang thuê, gia hạn, quá hạn hoặc chờ xử lý trả xe. Cần xác nhận xe đang được thuê ngoài hệ thống hay chỉ là dữ liệu mẫu. Sau khi xác nhận mới tạo hồ sơ thuê còn thiếu hoặc đưa xe về trạng thái sẵn sàng.'),
('Đơn đã hoàn tất thiếu dữ liệu đối soát', 'Có 1 đơn COMPLETED chưa có ngày trả thực tế; bảng payments hiện có 0 bản ghi. Vì thiếu chứng từ và ngày thực trả, báo cáo mới chưa đưa đơn này vào doanh thu theo ngày thực trả. Đề xuất đối chiếu biên nhận, lịch sử PayOS hoặc sổ thu tiền trước khi bổ sung; không tự giả lập thanh toán để làm đẹp báo cáo.'),
('Các bảng cần giữ lại', 'Giữ payments dù đã tích hợp PayOS: bảng này lưu khoản thu, liên kết nhà cung cấp, mã giao dịch, trạng thái đối soát và khoản hoàn. Giữ stations ở thời điểm này vì motorbikes còn quan hệ khóa ngoại; việc có một cửa hàng không làm quan hệ hiện tại tự mất tác dụng. Có thể đơn giản hóa sau bằng migration có kiểm tra phụ thuộc. Bảng deposit_attempts đang rỗng; chưa xóa khi chưa hoàn tất kiểm tra các luồng cũ và phụ thuộc. Không xóa các bảng khác chỉ vì chưa được giao diện mới sử dụng.'),
('Cấu hình cần kiểm tra trước khi vận hành', 'Kiểm tra địa chỉ, tọa độ, giờ mở cửa, phí giao nhận, tỷ lệ giữ xe, thời hạn giữ xe và điều khoản trong Cài đặt cửa hàng. Bản sao đang khởi tạo phần cấu hình còn thiếu từ nội dung giao diện cũ. Đăng nhập Google chưa có cấu hình OAuth hoạt động; đã bỏ nút giả chỉ hiển thị thông báo. Thanh toán PayOS, gửi email và OCR thật chưa được xác minh bằng giao dịch thực tế.'),
('Các thay đổi đã kiểm tra', '21 bài kiểm thử tự động đã đạt. Đã kiểm tra trên trình duyệt: lưu hồ sơ tiếng Việt, lịch sử thuê, admin lưu và tải lại thông tin xe trên MySQL, đăng xuất admin, đăng nhập staff và trang thông tin cửa hàng. Biên bản giao trả đọc lại dữ liệu đã lưu và chặn ghi đè thông tin lúc giao sau bàn giao. Đây chưa phải xác nhận toàn bộ dự án đã sẵn sàng vận hành.')]
for title,body in sections:
    doc.add_paragraph(title,style='Heading 1');doc.add_paragraph(body)
doc.add_paragraph('Thông tin cần chủ cửa hàng trả lời',style='Heading 1')
doc.add_paragraph('11 xe và các đơn cũ là dữ liệu thử hay dữ liệu vận hành thật? Với 3 xe đang mang trạng thái RENTED, hiện có khách đang giữ xe không? Nếu dữ liệu là thật, vui lòng đối chiếu biển số, ngày trả và chứng từ trước khi cho phép hiệu chỉnh.')
doc.save(out)
print(out)
